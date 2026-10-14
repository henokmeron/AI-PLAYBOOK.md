"""Bounded structured-data analysis over CSV and XLSX tables in the authorised root.

Operations: ``describe``, ``select`` (filter + project + limit) and ``aggregate``
(sum/mean/min/max/count, optionally grouped). Arithmetic uses Decimal on the
stored cell values so results match the calculator tool.
"""
from __future__ import annotations

import csv
import io
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path, PurePosixPath
from typing import Any

from ..tools import Tool, ToolInputError
from .files import AuthorisedRoot

MAX_TABLE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 20_000
MAX_COLUMNS = 100
MAX_CELL_CHARS = 1000
MAX_SELECT_ROWS = 50
MAX_GROUPS = 100
MAX_XLSX_UNCOMPRESSED = 60 * 1024 * 1024
MAX_XLSX_MEMBER = 30 * 1024 * 1024
FILTER_OPS = ["eq", "ne", "gt", "ge", "lt", "le", "contains"]
AGG_OPS = ["sum", "mean", "min", "max", "count"]


def parse_number(value: Any) -> Decimal | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip().replace(",", "")
    for prefix in ("£", "$", "€"):
        if text.startswith(prefix):
            text = text[len(prefix):]
        elif text.startswith("-" + prefix):
            text = "-" + text[len(prefix) + 1:]
    if not text:
        return None
    try:
        number = Decimal(text)
    except InvalidOperation:
        return None
    return number if number.is_finite() else None


def _plain(value: Decimal) -> str:
    text = format(value, "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text or "0"


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        value = Decimal(str(value))
    if isinstance(value, Decimal):
        return _plain(value)
    if hasattr(value, "isoformat"):
        return value.isoformat()
    text = str(value).strip()
    if len(text) > MAX_CELL_CHARS:
        raise ToolInputError("a cell exceeds the size limit")
    return text


def _from_rows(rows: list[list[str]]) -> tuple[list[str], list[list[str]]]:
    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        raise ToolInputError("table is empty")
    header = rows[0]
    if len(header) > MAX_COLUMNS:
        raise ToolInputError("table has too many columns")
    columns = []
    for index, name in enumerate(header):
        label = name or f"column_{index + 1}"
        while label in columns:
            label += "_dup"
        columns.append(label)
    body = [(row + [""] * len(columns))[:len(columns)] for row in rows[1:]]
    return columns, body


def load_csv(path: Path) -> tuple[list[str], list[list[str]]]:
    raw = path.read_bytes()
    if len(raw) > MAX_TABLE_BYTES:
        raise ToolInputError("table file exceeds the size limit")
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ToolInputError("CSV is not valid UTF-8") from None
    rows: list[list[str]] = []
    try:
        for row in csv.reader(io.StringIO(text, newline="")):
            if len(rows) > MAX_ROWS:
                raise ToolInputError("table has too many rows")
            if len(row) > MAX_COLUMNS:
                raise ToolInputError("table has too many columns")
            rows.append([_cell(c) for c in row])
    except csv.Error as exc:
        raise ToolInputError(f"invalid CSV: {exc}") from None
    return _from_rows(rows)


def _check_xlsx_container(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            total = 0
            for info in archive.infolist():
                member = PurePosixPath(info.filename)
                if member.is_absolute() or ".." in member.parts:
                    raise ToolInputError("unsafe entry in workbook")
                if info.file_size > MAX_XLSX_MEMBER:
                    raise ToolInputError("workbook component exceeds the size limit")
                total += info.file_size
                if total > MAX_XLSX_UNCOMPRESSED:
                    raise ToolInputError("workbook exceeds the uncompressed size limit")
    except zipfile.BadZipFile:
        raise ToolInputError("not a valid XLSX workbook") from None


def load_xlsx(path: Path, sheet: str | None) -> tuple[list[str], list[list[str]]]:
    if path.stat().st_size > MAX_TABLE_BYTES:
        raise ToolInputError("table file exceeds the size limit")
    _check_xlsx_container(path)
    try:
        import openpyxl
    except ImportError:
        raise ToolInputError("XLSX analysis requires openpyxl, which is not installed") from None
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if sheet is None:
            worksheet = workbook.worksheets[0]
        elif sheet in workbook.sheetnames:
            worksheet = workbook[sheet]
        else:
            raise ToolInputError(f"sheet {sheet!r} not found; available: {workbook.sheetnames[:20]}")
        rows: list[list[str]] = []
        for row in worksheet.iter_rows(values_only=True, max_col=MAX_COLUMNS + 1):
            if len(rows) > MAX_ROWS:
                raise ToolInputError("table has too many rows")
            if len(row) > MAX_COLUMNS and any(v is not None for v in row[MAX_COLUMNS:]):
                raise ToolInputError("table has too many columns")
            rows.append([_cell(v) for v in row[:MAX_COLUMNS]])
    finally:
        workbook.close()
    return _from_rows(rows)


def load_table(root: AuthorisedRoot, relative: str, sheet: str | None = None) -> tuple[list[str], list[list[str]]]:
    path = root.resolve(relative)
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        return load_csv(path)
    if suffix == ".xlsx":
        return load_xlsx(path, sheet)
    raise ToolInputError("only .csv and .xlsx tables are supported")


def _col(columns: list[str], name: str) -> int:
    if name not in columns:
        raise ToolInputError(f"unknown column {name!r}; columns are {columns[:MAX_COLUMNS]}")
    return columns.index(name)


def _matches(cell: str, op: str, value: Any) -> bool:
    if op == "contains":
        return str(value).casefold() in cell.casefold()
    left, right = parse_number(cell), parse_number(value)
    if left is not None and right is not None:
        a, b = left, right
    else:
        a, b = cell.casefold(), str(value).strip().casefold()
        if op not in {"eq", "ne"}:
            return False
    return {"eq": a == b, "ne": a != b, "gt": a > b, "ge": a >= b, "lt": a < b, "le": a <= b}[op]


def _filtered(columns: list[str], rows: list[list[str]], filters: list[dict[str, Any]]) -> list[list[str]]:
    resolved = [(_col(columns, f["column"]), f["op"], f["value"]) for f in filters]
    return [row for row in rows if all(_matches(row[i], op, v) for i, op, v in resolved)]


def _aggregate(values: list[str], op: str) -> dict[str, Any]:
    if op == "count":
        return {"value": str(sum(1 for v in values if v != "")), "count": len(values)}
    numbers = [n for n in (parse_number(v) for v in values) if n is not None]
    skipped = len(values) - len(numbers)
    if not numbers:
        return {"value": None, "count": 0, "skipped_non_numeric": skipped}
    if op == "sum":
        result = sum(numbers, Decimal(0))
    elif op == "mean":
        result = sum(numbers, Decimal(0)) / Decimal(len(numbers))
    elif op == "min":
        result = min(numbers)
    else:
        result = max(numbers)
    return {"value": _plain(result), "count": len(numbers), "skipped_non_numeric": skipped}


def analyse(root: AuthorisedRoot, arguments: dict[str, Any]) -> dict[str, Any]:
    columns, rows = load_table(root, arguments["path"], arguments.get("sheet"))
    operation = arguments["operation"]
    rows = _filtered(columns, rows, arguments.get("filters", []))
    out: dict[str, Any] = {"path": arguments["path"], "operation": operation,
                           "columns": columns, "matched_rows": len(rows), "untrusted": True}
    if operation == "describe":
        numeric = [c for i, c in enumerate(columns)
                   if rows and sum(parse_number(r[i]) is not None for r in rows) >= max(1, len(rows) // 2)]
        out["numeric_columns"] = numeric
        out["sample"] = rows[:5]
        return out
    if operation == "select":
        wanted = arguments.get("columns") or columns
        indexes = [_col(columns, name) for name in wanted]
        limit = arguments.get("limit", 20)
        out["rows"] = [[row[i] for i in indexes] for row in rows[:limit]]
        out["selected_columns"] = wanted
        out["truncated"] = len(rows) > limit
        return out
    target = _col(columns, arguments["column"]) if arguments.get("column") else None
    agg = arguments.get("aggregate", "count")
    if target is None and agg != "count":
        raise ToolInputError("aggregate requires a column unless aggregate is count")
    def values_for(group_rows: list[list[str]]) -> list[str]:
        return [r[target] for r in group_rows] if target is not None else ["x"] * len(group_rows)
    group_by = arguments.get("group_by")
    if group_by:
        gi = _col(columns, group_by)
        groups: dict[str, list[list[str]]] = {}
        for row in rows:
            groups.setdefault(row[gi], []).append(row)
            if len(groups) > MAX_GROUPS:
                raise ToolInputError("too many groups")
        out["groups"] = [{"key": key, **_aggregate(values_for(g), agg)} for key, g in groups.items()]
    else:
        out["result"] = _aggregate(values_for(rows), agg)
    out["aggregate"] = agg
    return out


def table_analysis_tool(root: AuthorisedRoot) -> Tool:
    return Tool(
        name="table_analysis",
        description=(
            "Analyse a CSV/XLSX table in the authorised folder: describe columns, select filtered "
            "rows, or aggregate (sum/mean/min/max/count) a column, optionally grouped. Exact decimals."
        ),
        input_schema={
            "type": "object",
            "properties": {
                "path": {"type": "string", "minLength": 1, "maxLength": 255},
                "sheet": {"type": "string", "maxLength": 64},
                "operation": {"enum": ["describe", "select", "aggregate"]},
                "columns": {"type": "array", "maxItems": MAX_COLUMNS, "items": {"type": "string", "maxLength": 200}},
                "filters": {"type": "array", "maxItems": 10, "items": {
                    "type": "object",
                    "properties": {
                        "column": {"type": "string", "maxLength": 200},
                        "op": {"enum": FILTER_OPS},
                        "value": {"type": ["string", "number"]},
                    },
                    "required": ["column", "op", "value"],
                    "additionalProperties": False,
                }},
                "column": {"type": "string", "maxLength": 200},
                "aggregate": {"enum": AGG_OPS},
                "group_by": {"type": "string", "maxLength": 200},
                "limit": {"type": "integer", "minimum": 1, "maximum": MAX_SELECT_ROWS},
            },
            "required": ["path", "operation"],
            "additionalProperties": False,
        },
        output_schema={
            "type": "object",
            "properties": {
                "path": {"type": "string"}, "operation": {"type": "string"},
                "columns": {"type": "array"}, "matched_rows": {"type": "integer"},
                "untrusted": {"const": True},
            },
            "required": ["path", "operation", "columns", "matched_rows", "untrusted"],
        },
        handler=lambda arguments: analyse(root, arguments),
        timeout_seconds=30.0,
        tags=("read", "data"),
    )
